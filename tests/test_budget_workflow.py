import tempfile
import unittest
import os
from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path
from docx.oxml.ns import qn
from unittest.mock import patch

from core.env_loader import load_local_env

load_local_env()

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
from sqlalchemy import JSON, MetaData, create_engine
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from crud import budget as budgets
from crud import proforma as proformas
from crud import site as sites
from database.base import Base
from database.models import Budget, Offre
from database.schemas import BudgetCreate, ProformaCreate, SiteCreate
from services import budget_documents, proforma_pdf, proforma_word


class BudgetWorkflowTests(unittest.TestCase):
    def setUp(self):
        # Exercise the ORM against a disposable database; never use DATABASE_URL.
        self.engine = create_engine("sqlite://")
        metadata = MetaData()
        for table in Base.metadata.tables.values():
            table.to_metadata(metadata)
        for table in metadata.tables.values():
            for column in table.columns:
                if isinstance(column.type, JSONB):
                    column.type = JSON()
                    column.server_default = None
        names = ("offre", "budgets", "proformas", "site", "demandes_team_building")
        metadata.create_all(self.engine, tables=[metadata.tables[name] for name in names])
        self.db = Session(self.engine)
        self.db.add(Offre(id=1, demande_id=10, titre="Seminaire", montant_total=0))
        self.db.add(Offre(id=2, demande_id=20, titre="Autre offre", montant_total=0))
        self.db.commit()
        self.output = tempfile.TemporaryDirectory()
        self.addCleanup(self.output.cleanup)
        self.addCleanup(self.engine.dispose)
        self.addCleanup(self.db.close)
        patcher = patch.object(budget_documents, "DEFAULT_OUTPUT_DIR", Path(self.output.name))
        patcher.start()
        self.addCleanup(patcher.stop)

    def payload(self, site_id=1, **overrides):
        values = {
            "offre_id": 1, "demande_team_building_id": 10, "site_id": site_id,
            "titre": "Seminaire", "client": "Client test", "nombre_personnes": 13,
            "date_budget": date(2026, 9, 9), "duree_jours": 2,
            "sections": [{"nom": "Transport", "prestations": [{
                "designation": "Abidjan - Bouake", "nombre_jours": 2,
                "quantite": 1, "prix_unitaire": 300000,
            }]}],
            "frais_agence": 12345, "taux_tva_frais_agence": 0,
        }
        values.update(overrides)
        return BudgetCreate(**values)

    def create(self, site_id=1, **overrides):
        return budgets.create_budget(self.db, self.payload(site_id, **overrides))

    def test_single_approval_and_explicit_replacement(self):
        first, second = self.create(), self.create(2)
        budgets.validate_budget(self.db, first)
        with self.assertRaises(budgets.BudgetConflict):
            budgets.validate_budget(self.db, second)
        with self.assertRaises(budgets.BudgetConflict):
            budgets.validate_budget(self.db, second, remplacer_budget_id=999)
        budgets.validate_budget(self.db, second, remplacer_budget_id=first.id)
        self.assertEqual(first.statut, "genere")
        self.assertEqual(second.statut, "valide")
        self.assertEqual(self.db.query(Budget).filter_by(statut="valide").count(), 1)
        self.assertEqual(self.db.get(Offre, 1).montant_total, second.total_ttc)

    def test_drink_fees_persist_from_site_to_budget_exports_and_proforma(self):
        tarifs = {
            "droit_bouchon_vin": 3000,
            "droit_bouchon_champagne": 12000,
            "droit_bouchon_alcool_fort": 7000,
            "frais_soft": 1500,
        }
        site = sites.create_site(self.db, SiteCreate(
            nom_site="Site boissons", a_restauration=True, tarifs_restauration=tarifs,
        ))
        self.db.expire_all()
        self.assertEqual(sites.get_site(self.db, site.id_site).tarifs_restauration, tarifs)
        quantities = {
            "droit_bouchon_vin": 6,
            "droit_bouchon_champagne": 2,
            "droit_bouchon_alcool_fort": 3,
        }
        labels = [
            ("droit_bouchon_vin", "Droit de bouchon vin (par bouteille)", 6),
            ("droit_bouchon_champagne", "Droit de bouchon champagne (par bouteille)", 2),
            ("droit_bouchon_alcool_fort", "Droit de bouchon alcools forts (par bouteille)", 3),
            ("frais_soft", "Frais soft (par personne)", 40),
        ]
        budget = self.create(
            site_id=site.id_site, nombre_personnes=40, duree_jours=3, frais_agence=0,
            sections=[{"nom": "Restauration", "prestations": [
                {"designation": label, "nombre_jours": 1, "quantite": quantity,
                 "prix_unitaire": site.tarifs_restauration[key]}
                for key, label, quantity in labels
            ]}],
            options_selectionnees={"restauration_quantites": quantities},
        )
        self.db.expire_all()
        self.assertEqual(budget.options_selectionnees["restauration_quantites"], quantities)
        self.assertEqual(budget.total_ttc, Decimal("123000"))
        data = budgets.budget_document_data(budget)
        pdf = budget_documents.generate_budget_pdf(data, output_dir=self.output.name)
        self.assertTrue(Path(pdf).is_file())
        workbook = load_workbook(budget_documents.generate_budget_excel(data, output_dir=self.output.name))
        self.addCleanup(workbook.close)
        sheet = workbook.active
        for key, label, quantity in labels:
            row = next(cell.row for cell in sheet["B"] if cell.value == label)
            self.assertEqual(sheet.cell(row, 3).value, quantity)
            self.assertEqual(sheet.cell(row, 4).value, tarifs[key])
            self.assertEqual(sheet.cell(row, 5).value, 1)
        self.assertFalse(any(
            str(cell.value).startswith("Droit de bouchon vin :")
            for row in sheet for cell in row if cell.value
        ))
        budgets.validate_budget(self.db, budget)
        proforma = proformas.create_proforma(self.db, ProformaCreate(
            **proformas.build_proforma_from_budget(self.db, budget.id)
        ))
        self.assertEqual(proforma.total_ttc, Decimal("123000"))
        self.assertEqual([line["quantite"] for line in proforma.sections[0]["prestations"]], ["6", "2", "3", "40"])

    def test_database_rejects_two_approvals_even_without_crud(self):
        first, second = self.create(), self.create(2)
        budgets.validate_budget(self.db, first)
        second.statut = "valide"
        with self.assertRaises(IntegrityError):
            self.db.commit()
        self.db.rollback()
        self.assertEqual(self.db.query(Budget).filter_by(statut="valide").count(), 1)

    def test_create_and_update_cannot_bypass_validation(self):
        with self.assertRaises(ValueError):
            self.create(statut="valide")
        budget = self.create()
        with self.assertRaises(ValueError):
            budgets.update_budget(self.db, budget, {"statut": "valide"})

    def test_new_draft_does_not_change_offer_amount(self):
        offer = self.db.get(Offre, 1)
        offer.montant_total = 777
        self.db.commit()
        self.create()
        self.assertEqual(offer.montant_total, 777)

    def test_editing_approved_budget_requires_validation_again(self):
        budget = self.create()
        budgets.validate_budget(self.db, budget)
        budgets.update_budget(self.db, budget, {"frais_agence": 22222})
        self.assertEqual(budget.statut, "brouillon")
        self.assertEqual(self.db.get(Offre, 1).montant_total, 0)
        with self.assertRaises(ValueError):
            proformas.build_proforma_from_budget(self.db, budget.id)

    def test_proforma_prefill_snapshot_and_zero_vat(self):
        first, second = self.create(), self.create(2)
        budgets.validate_budget(self.db, first)
        values = proformas.build_proforma_from_budget(self.db, first.id)
        proforma = proformas.create_proforma(self.db, ProformaCreate(**values))
        self.assertEqual(proforma.budget_id, first.id)
        self.assertEqual(proforma.total_ttc, first.total_ttc)
        self.assertEqual(proforma.frais_agence, Decimal("12345"))
        self.assertEqual(proforma.tva_frais_agence, 0)
        snapshot = deepcopy(proforma.budget_snapshot)
        budgets.validate_budget(self.db, second, remplacer_budget_id=first.id)
        with self.assertRaises(ValueError):
            proformas.create_proforma(self.db, ProformaCreate(**values))
        proformas.update_proforma(self.db, proforma, {"objet": "Objet ajuste"})
        budgets.update_budget(self.db, first, {"frais_agence": 99999})
        self.db.refresh(proforma)
        self.assertEqual(proforma.budget_snapshot, snapshot)
        self.assertEqual(proforma.total_ttc, Decimal("612345"))
        with self.assertRaises(ValueError):
            proformas.update_proforma(self.db, proforma, {"budget_id": None})

    def test_proforma_rejects_unapproved_or_unrelated_budget(self):
        budget = self.create()
        with self.assertRaises(ValueError):
            proformas.build_proforma_from_budget(self.db, budget.id)
        budgets.validate_budget(self.db, budget)
        values = proformas.build_proforma_from_budget(self.db, budget.id)
        values["offre_id"] = 2
        with self.assertRaises(ValueError):
            proformas.create_proforma(self.db, ProformaCreate(**values))
        values["budget_id"] = None
        with self.assertRaises(ValueError):
            proformas.create_proforma(self.db, ProformaCreate(**values))

    def test_group_excel_and_individual_pdf_regeneration(self):
        members = budgets.create_multi_site_budgets(self.db, [self.payload(), self.payload(2)])
        self.assertEqual(members[0].groupe_reference, members[1].groupe_reference)
        self.assertEqual(members[0].fichier_excel, members[1].fichier_excel)
        budgets.validate_budget(self.db, members[1])
        budgets.generate_documents_for_budget(self.db, members[0])
        workbook = load_workbook(budgets.get_budget_excel_path(members[0]))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Comparatif", "Seminaire", "Seminaire (2)"])
        self.assertEqual(workbook["Comparatif"]["C3"].value, "Validé")
        self.assertTrue(workbook["Comparatif"]["G3"].value.startswith("='Seminaire (2)'!F"))
        self.assertTrue(budgets.get_budget_pdf_path(members[0]).is_file())
        self.assertEqual(members[0].fichier_excel, members[1].fichier_excel)

    def test_group_is_atomic_and_rejects_duplicate_sites_or_mixed_offers(self):
        with self.assertRaises(ValueError):
            budgets.create_multi_site_budgets(self.db, [self.payload(), self.payload()])
        with self.assertRaises(ValueError):
            budgets.create_multi_site_budgets(self.db, [self.payload(), self.payload(2, offre_id=2)])
        with self.assertRaises(ValueError):
            budgets.create_multi_site_budgets(self.db, [self.payload(), self.payload(2, sections=[])])
        self.assertEqual(self.db.query(Budget).count(), 0)

    def test_pdf_word_preserve_budget_fees_and_vat(self):
        budget = self.create(taux_tva_frais_agence=18)
        budgets.validate_budget(self.db, budget)
        values = proformas.build_proforma_from_budget(self.db, budget.id)
        proforma = proformas.create_proforma(self.db, ProformaCreate(**values))
        self.assertEqual(proforma.total_ttc, budget.total_ttc)
        data = proformas._document_data(proforma)
        for module, generator in ((proforma_pdf, proforma_pdf.generate_proforma_pdf), (proforma_word, proforma_word.generate_proforma_word)):
            with patch.object(module, "calculate_totals", wraps=module.calculate_totals) as calculate:
                path = generator(data, output_dir=self.output.name)
                self.assertTrue(Path(path).is_file())
                self.assertIsNone(calculate.call_args.kwargs["agency_fee_rate"])
                self.assertEqual(calculate.call_args.args[1], Decimal("12345"))
        document = Document(Path(self.output.name) / f"{proforma.reference}.docx")
        text_content = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("12 345", text_content)
        self.assertIn("2 222", text_content)

    def test_percentage_mode_is_calculated_by_server_and_copied_to_proforma(self):
        budget = self.create(mode_frais_agence="pourcentage", frais_agence=1, taux_tva_frais_agence=18)
        self.assertEqual(budget.frais_agence, Decimal("105000"))
        self.assertEqual(budget.total_ttc, Decimal("723900"))
        budgets.validate_budget(self.db, budget)
        values = proformas.build_proforma_from_budget(self.db, budget.id)
        self.assertEqual(values["mode_frais_agence"], "pourcentage")
        proforma = proformas.create_proforma(self.db, ProformaCreate(**values))
        self.assertEqual(proforma.total_ttc, budget.total_ttc)
        updated = proformas.update_proforma(self.db, proforma, {"mode_frais_agence": "montant", "frais_agence": 0})
        self.assertEqual(updated.frais_agence, 0)
        self.assertEqual(updated.total_ttc, Decimal("600000"))
        updated = proformas.update_proforma(self.db, proforma, {"mode_frais_agence": "pourcentage", "frais_agence": 1})
        self.assertEqual(updated.total_ttc, Decimal("723900"))

    def test_word_stamp_is_only_in_ivoir_trips_signature_cell(self):
        budget = self.create(mode_frais_agence="pourcentage")
        budgets.validate_budget(self.db, budget)
        proforma = proformas.create_proforma(self.db, ProformaCreate(**proformas.build_proforma_from_budget(self.db, budget.id)))
        path = proforma_word.generate_proforma_word(proformas._document_data(proforma), output_dir=self.output.name)
        document = Document(path)
        signature = next(table for table in document.tables if 'POUR LE CLIENT' in table.cell(0, 0).text)
        self.assertFalse(signature.cell(0, 0)._tc.xpath('.//pic:pic'))
        drawings = signature.cell(0, 1)._tc.xpath('.//a:blip')
        self.assertEqual(len(drawings), 1)
        relationship = document.part.related_parts[drawings[0].get(qn('r:embed'))]
        self.assertEqual(relationship.blob, (proforma_word.ASSETS_DIR / 'signature1.png').read_bytes())

        services = next(table for table in document.tables if table.cell(0, 0).text == "DESIGNATION")
        for cell in services.rows[0].cells:
            shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            self.assertEqual(shading.get(qn("w:fill")), "C74C0A")
            self.assertEqual(cell.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            header_run = next(run for run in cell.paragraphs[0].runs if run.text)
            self.assertEqual(header_run.font.size.pt, 8.5)
        self.assertEqual(services.rows[2].cells[0].paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        for cell in services.rows[2].cells[1:]:
            self.assertEqual(cell.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(services.rows[3].cells[0].paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(services.rows[3].cells[4].paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_pdf_contains_ivoir_trips_stamp(self):
        budget = self.create(mode_frais_agence="pourcentage")
        budgets.validate_budget(self.db, budget)
        proforma = proformas.create_proforma(self.db, ProformaCreate(**proformas.build_proforma_from_budget(self.db, budget.id)))
        path = Path(proforma_pdf.generate_proforma_pdf(proformas._document_data(proforma), output_dir=self.output.name))
        signature = proforma_pdf._signature_image()
        self.assertIsNotNone(signature)
        self.assertEqual(Path(signature.filename).name, "signature1.png")
        self.assertEqual(signature.hAlign, "CENTER")
        self.assertGreaterEqual(path.read_bytes().count(b"/Subtype /Image"), 3)
        self.assertTrue(proforma_pdf.is_proforma_pdf_current(path))
        os.utime(path, (1, 1))
        self.assertFalse(proforma_pdf.is_proforma_pdf_current(path))
        styles = proforma_pdf._styles()
        self.assertEqual(proforma_pdf.IVT_TABLE_HEADER.hexval(), "0xc74c0a")
        self.assertEqual(styles["table_header"].fontSize, 8.5)
        self.assertEqual(styles["table_cell"].alignment, 1)

    def test_percentage_excel_uses_formula_and_manual_amount_remains_literal(self):
        for mode in ("pourcentage", "montant"):
            data = {**self.payload().model_dump(), "reference": mode, "mode_frais_agence": mode}
            path = budget_documents.generate_budget_excel(data, output_dir=self.output.name)
            workbook = load_workbook(path)
            sheet = workbook.active
            agency_row = next(row[0].row + 1 for row in sheet.iter_rows() if any(str(cell.value).startswith("HONORAIRES D'AGENCE") for cell in row))
            value = sheet.cell(agency_row, 6).value
            if mode == "pourcentage":
                self.assertIn("*17.5%", value)
            else:
                self.assertEqual(value, 12345)
            workbook.close()

    def test_old_word_is_regenerated_after_template_update(self):
        budget = self.create()
        budgets.validate_budget(self.db, budget)
        proforma = proformas.create_proforma(self.db, ProformaCreate(**proformas.build_proforma_from_budget(self.db, budget.id)))
        path = Path(proforma_word.generate_proforma_word(proformas._document_data(proforma), output_dir=self.output.name))
        with patch.object(proformas, "get_proforma_word_path", return_value=path):
            self.assertEqual(proformas.get_word_path(proforma), path)
            os.utime(path, (1, 1))
            self.assertIsNone(proformas.get_word_path(proforma))
            proforma_word.generate_proforma_word(proformas._document_data(proforma), output_dir=self.output.name)
            self.assertEqual(proformas.get_word_path(proforma), path)

    def test_pdf_word_honor_explicit_fee_modes_without_budget(self):
        data = {
            "reference": "PRO-FEE-MODE", "client": "Client test", "nombre_personnes": 13,
            "objet": "Seminaire", "date_proforma": date(2026, 9, 9), "pole": "teambuilding",
            "sections": self.payload().model_dump()["sections"], "taux_tva_frais_agence": 0,
        }
        for mode, fees in (("montant", 0), ("montant", 12345), ("pourcentage", 1)):
            data.update(mode_frais_agence=mode, frais_agence=fees)
            for module, generator in ((proforma_pdf, proforma_pdf.generate_proforma_pdf), (proforma_word, proforma_word.generate_proforma_word)):
                with self.subTest(mode=mode, fees=fees, generator=generator.__name__):
                    with patch.object(module, "calculate_totals", wraps=module.calculate_totals) as calculate:
                        generator(data, output_dir=self.output.name)
                    expected = Decimal("0.175") if mode == "pourcentage" else None
                    self.assertEqual(calculate.call_args.kwargs["agency_fee_rate"], expected)
                    self.assertEqual(calculate.call_args.args[1], fees)


if __name__ == "__main__":
    unittest.main()
