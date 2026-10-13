import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from services import budget_documents as documents


class BudgetTemplateTests(unittest.TestCase):
    def setUp(self):
        self.output = tempfile.TemporaryDirectory()
        self.addCleanup(self.output.cleanup)
        self.data = {
            "reference": "BUD-TEMPLATE", "offre_id": 1, "site_nom": "Site / Assinie",
            "client": "Entreprise exemple", "titre": "Séminaire team building",
            "date_budget": date(2026, 10, 13), "nombre_personnes": 18,
            "duree_jours": 2, "nombre_nuitees": 1,
            "sections": [{"nom": "Transport", "prestations": [
                {"designation": "Car", "quantite": 1, "nombre_jours": 2, "prix_unitaire": 170000},
                {"designation": "Montant ajusté", "quantite": 2, "nombre_jours": 1,
                 "prix_unitaire": 1000, "montant_ht": 1500},
            ]}],
            "frais_agence": 2000000, "taux_tva_frais_agence": 18,
            "mode_frais_agence": "montant",
            "taxes_selectionnees": [
                {"code": "TDT", "nom": "Taxe touristique", "taux": 2.5, "montant": 83250.25},
                {"code": "NUITEE", "nom": "Taxe de nuitée", "tarif": 2000,
                 "nombre_chambres": 18, "nombre_nuitees": 1, "montant": 36000},
            ],
        }

    def workbook(self, **changes):
        workbook = load_workbook(documents.generate_budget_excel(
            {**self.data, **changes}, self.output.name,
        ))
        self.addCleanup(workbook.close)
        return workbook

    @staticmethod
    def find(sheet, label):
        return next(cell for row in sheet for cell in row if cell.value == label)

    def test_reference_layout_is_printable_and_amount_override_survives(self):
        sheet = self.workbook().active
        self.assertEqual(sheet["B7"].value, "BUDGET ESTIMATIF")
        self.assertEqual(sheet["B8"].value, "SÉMINAIRE TEAM BUILDING")
        self.assertEqual(sheet["C11"].value, 18)
        self.assertEqual(sheet["F10"].value, 2)
        self.assertEqual(sheet["F11"].value, 1)
        self.assertEqual([sheet.cell(15, c).value for c in range(2, 7)],
                         ["Désignation", "Quantité", "Coût unitaire", "Nombre de jours", "Coût total"])
        self.assertEqual(sheet["B16"].fill.fgColor.rgb[-6:], "FFC000")
        self.assertIn("B16:F16", sheet.merged_cells)
        adjusted = self.find(sheet, "Montant ajusté").row
        self.assertEqual(sheet.cell(adjusted, 6).value, 1500)
        self.assertEqual(self.find(sheet, "SOUS-TOTAL TRANSPORT").fill.fgColor.rgb[-6:], "AEAAAA")
        notice = self.find(sheet, documents.BUDGET_DISCLAIMER)
        self.assertIn(str(notice.row), sheet.print_area)
        self.assertEqual(sheet.page_setup.orientation, "portrait")
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(len(sheet._images), 1)

    def test_taxes_are_displayed_once_and_not_added_to_percentage_fee_base(self):
        for mode in ("montant", "pourcentage"):
            with self.subTest(mode=mode):
                sheet = self.workbook(mode_frais_agence=mode).active
                tax_row = self.find(sheet, "Taxe touristique (2.5 %)").row
                self.assertEqual(sheet.cell(tax_row, 6).value, 83250.25)
                self.assertIn(".00", sheet.cell(tax_row, 6).number_format)
                night_row = self.find(sheet, "Taxe de nuitée").row
                self.assertEqual([sheet.cell(night_row, c).value for c in range(3, 7)], [18, 2000, 1, 36000])
                tax_subtotal = self.find(sheet, "SOUS-TOTAL TAXES ET CONTRIBUTIONS").row
                implementation = self.find(sheet, "SOUS-TOTAL MISE EN ŒUVRE").row
                self.assertIn(f"F{tax_subtotal}", sheet.cell(implementation, 6).value)
                agency_row = self.find(sheet, documents.AGENCY_DETAILS).row
                if mode == "pourcentage":
                    self.assertNotIn(f"F{tax_subtotal}", sheet.cell(agency_row, 6).value)
                    self.assertIn("*17.5%", sheet.cell(agency_row, 6).value)
                else:
                    self.assertEqual(sheet.cell(agency_row, 6).value, 2000000)
                total_row = self.find(sheet, "TOTAL TTC").row
                before_vat = self.find(sheet, "TOTAL AVANT TVA SUR HONORAIRES").row
                vat_row = self.find(sheet, "TVA 18 %").row
                self.assertEqual(sheet.cell(total_row, 6).value, f"=F{before_vat}+F{vat_row}")
                per_person = self.find(sheet, "TOTAL TTC / PERSONNE").row
                self.assertEqual(sheet.cell(per_person, 6).value, f"=F{total_row}/C11")
        totals = documents._document_totals(self.data)
        self.assertEqual(totals["total_ttc"], Decimal("2820750.25"))

    def test_multi_site_links_follow_dynamic_totals_and_duplicate_names(self):
        path = documents.generate_multi_site_budget_excel(
            [self.data, {**self.data, "reference": "BUD-SECOND", "taxes_selectionnees": []}],
            "GROUP-TEMPLATE", self.output.name,
        )
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(len(set(workbook.sheetnames)), 3)
        for row, sheet in enumerate(workbook.worksheets[1:], 2):
            total_row = self.find(sheet, "TOTAL TTC").row
            self.assertEqual(workbook["Comparatif"].cell(row, 7).value, f"='{sheet.title}'!F{total_row}")

    def test_pdf_generation_handles_many_sections_and_old_template_detection(self):
        data = {**self.data, "sections": self.data["sections"] * 15}
        path = Path(documents.generate_budget_pdf(data, self.output.name))
        self.assertTrue(path.read_bytes().startswith(b"%PDF"))
        self.assertTrue(documents.is_budget_pdf_current(path))
        os.utime(path, (1, 1))
        self.assertFalse(documents.is_budget_pdf_current(path))

    def test_zero_participants_cannot_create_division_by_zero_formula(self):
        for generate in (documents.generate_budget_excel, documents.generate_budget_pdf):
            with self.assertRaises(ValueError):
                generate({**self.data, "nombre_personnes": 0}, self.output.name)


if __name__ == "__main__":
    unittest.main()
