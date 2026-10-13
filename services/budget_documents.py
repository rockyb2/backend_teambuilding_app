from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import quote_sheetname
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from services.agency_fees import resolve_agency_fee_rate
from services.taxes import add_persisted_taxes


BASE_DIR = Path(__file__).resolve().parents[1]
ASSETS_DIR = BASE_DIR / "assets" / "proforma"
DEFAULT_OUTPUT_DIR = BASE_DIR / "uploads" / "budgets"


def _decimal(value: Any, field_name: str = "montant") -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, str):
        value = (
            value.upper()
            .replace("F CFA", "")
            .replace("FCFA", "")
            .replace("XOF", "")
            .replace("\u202f", "")
            .replace(" ", "")
            .replace(",", ".")
        )
    try:
        amount = Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name} doit être un nombre valide.") from exc
    if amount < 0:
        raise ValueError(f"{field_name} ne peut pas être négatif.")
    return amount


def _quantity(value: Any, field_name: str = "quantité") -> Decimal:
    if value in (None, ""):
        return Decimal("1")
    try:
        quantity = Decimal(str(value).replace(",", "."))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name} doit être un nombre valide.") from exc
    if quantity < 0:
        raise ValueError(f"{field_name} ne peut pas être négative.")
    return quantity


def _format_quantity(value: Any) -> str:
    if value in (None, ""):
        return ""
    try:
        decimal_value = Decimal(str(value))
    except InvalidOperation:
        return str(value)
    return format(decimal_value.normalize(), "f")


def _format_fcfa(value: Any) -> str:
    amount = int(_decimal(value))
    return f"{amount:,}".replace(",", " ") + " F CFA"


def _display_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value).strftime("%d/%m/%Y")
        except ValueError:
            return value
    return str(value or "")


def _excel_number(value: Any, default: int | float = 0) -> int | float:
    if value in (None, ""):
        return default
    if isinstance(value, str):
        value = (
            value.upper()
            .replace("F CFA", "")
            .replace("FCFA", "")
            .replace("XOF", "")
            .replace("\u202f", "")
            .replace(" ", "")
            .replace(",", ".")
        )
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return default
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _safe_sheet_title(value: Any, fallback: str = "Budget estimatif") -> str:
    title = str(value or fallback).strip() or fallback
    for forbidden in ("\\", "/", "?", "*", "[", "]", ":"):
        title = title.replace(forbidden, " ")
    title = " ".join(title.split()).strip("'")
    return title[:31] or fallback[:31]


def _budget_location(data: dict[str, Any]) -> str:
    location = (
        data.get("lieu")
        or data.get("site_localisation")
        or data.get("site_nom")
        or data.get("destination")
        or "-"
    )
    return str(location)


def _budget_nights(data: dict[str, Any]) -> int:
    explicit_value = data.get("nombre_nuitees", data.get("nuits"))
    if explicit_value not in (None, ""):
        return int(_excel_number(explicit_value, 0))
    days = int(_excel_number(data.get("duree_jours"), 1) or 1)
    return max(days - 1, 0)


def _budget_notes(data: dict[str, Any]) -> list[str]:
    notes = str(data.get("notes") or "").strip()
    if notes:
        return [line.strip() for line in notes.splitlines() if line.strip()]
    return [
        "Jus/Sucrerie : 2 canettes par personne",
        "Bière : 2 bouteilles par personne",
        "Vin : 1 bouteille pour 4 personnes",
    ]


def _safe_document_path(reference: str, extension: str, output_dir: str | Path | None = None) -> Path:
    clean_reference = "".join(
        char for char in str(reference or "budget") if char.isalnum() or char in ("-", "_")
    ).strip("-_")
    if not clean_reference:
        clean_reference = "budget"
    directory = Path(output_dir) if output_dir else DEFAULT_OUTPUT_DIR
    directory.mkdir(parents=True, exist_ok=True)
    return (directory / f"{clean_reference}.{extension.lstrip('.')}").resolve()


def _first_existing_asset(*filenames: str) -> Path | None:
    for filename in filenames:
        path = ASSETS_DIR / filename
        if path.exists():
            return path
    return None


def _image_dimensions(path: Path, max_width: float, max_height: float) -> tuple[float, float]:
    image_width, image_height = ImageReader(str(path)).getSize()
    if not image_width or not image_height:
        return max_width, max_height
    scale = min(max_width / image_width, max_height / image_height)
    return image_width * scale, image_height * scale


def _paragraph(text: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(str(text or "")), style)


def _styles() -> dict[str, ParagraphStyle]:
    sample = getSampleStyleSheet()
    return {
        "normal": ParagraphStyle(
            "NormalBudget",
            parent=sample["Normal"],
            fontName="Helvetica",
            fontSize=8.2,
            leading=10.2,
            alignment=TA_LEFT,
        ),
        "small": ParagraphStyle(
            "SmallBudget",
            parent=sample["Normal"],
            fontName="Helvetica",
            fontSize=7.4,
            leading=9,
        ),
        "title": ParagraphStyle(
            "TitleBudget",
            parent=sample["Normal"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=17,
            alignment=TA_CENTER,
        ),
        "center": ParagraphStyle(
            "CenterBudget",
            parent=sample["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
        ),
        "section": ParagraphStyle(
            "SectionBudget",
            parent=sample["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.2,
            leading=10,
        ),
        "right": ParagraphStyle(
            "RightBudget",
            parent=sample["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            alignment=TA_RIGHT,
        ),
        "right_bold": ParagraphStyle(
            "RightBoldBudget",
            parent=sample["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.2,
            leading=10,
            alignment=TA_RIGHT,
        ),
    }


def normalize_budget_sections(sections: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], Decimal]:
    normalized_sections: list[dict[str, Any]] = []
    implementation_total = Decimal("0")

    for section_index, section in enumerate(sections or [], start=1):
        if not isinstance(section, dict):
            raise ValueError(f"La section {section_index} doit être un dictionnaire.")

        section_name = str(section.get("nom") or section.get("name") or "Prestations").strip()
        items = section.get("prestations") or section.get("lignes") or []
        if not isinstance(items, list):
            raise ValueError(f"Les lignes de la section '{section_name}' doivent former une liste.")

        normalized_items: list[dict[str, Any]] = []
        section_total = Decimal("0")
        for item_index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                raise ValueError(
                    f"La ligne {item_index} de la section '{section_name}' doit être un dictionnaire."
                )

            designation = str(item.get("designation") or item.get("libelle") or "").strip()
            if not designation:
                continue

            days = _quantity(item.get("nombre_jours", item.get("duree_jours", 1)), "nombre_jours")
            quantity = _quantity(item.get("quantite", 1), "quantite")
            unit_price = _decimal(item.get("prix_unitaire"), "prix_unitaire")
            explicit_amount = item.get("montant_ht")

            if explicit_amount not in (None, ""):
                line_total = _decimal(explicit_amount, "montant_ht")
            else:
                line_total = (days * quantity * unit_price).quantize(Decimal("1"), rounding=ROUND_HALF_UP)

            normalized_items.append(
                {
                    "designation": designation,
                    "nombre_jours": _format_quantity(days),
                    "quantite": _format_quantity(quantity),
                    "prix_unitaire": int(unit_price),
                    "montant_ht": int(line_total),
                    "source": str(item.get("source") or "").strip(),
                    "prestation_hoteliere": item.get("prestation_hoteliere") is True,
                    "prix_est_ttc": item.get("prix_est_ttc") is True,
                    "montant_saisi": item.get("montant_saisi"),
                    "prix_unitaire_saisi": item.get("prix_unitaire_saisi"),
                }
            )
            section_total += line_total

        if normalized_items:
            normalized_sections.append(
                {
                    "nom": section_name,
                    "prestations": normalized_items,
                    "sous_total": int(section_total),
                }
            )
            implementation_total += section_total

    return normalized_sections, implementation_total


def calculate_budget_totals(
    sections: list[dict[str, Any]],
    frais_agence: Any = 0,
    taux_tva_frais_agence: Any = 0,
    mode_frais_agence: str = "montant",
) -> dict[str, Any]:
    normalized_sections, implementation_total = normalize_budget_sections(sections)
    agency_rate = resolve_agency_fee_rate(mode_frais_agence)
    agency_fees = (
        (implementation_total * agency_rate).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        if agency_rate is not None else _decimal(frais_agence, "frais_agence")
    )
    try:
        vat_rate = Decimal(str(taux_tva_frais_agence or 0).replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError("Le taux de TVA doit être un nombre valide.") from exc
    if vat_rate < 0:
        raise ValueError("Le taux de TVA ne peut pas être négatif.")

    vat = (agency_fees * vat_rate / Decimal("100")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    total_ht = implementation_total + agency_fees
    total = total_ht + vat
    return {
        "sections": normalized_sections,
        "sous_total_ht": implementation_total,
        "frais_agence": agency_fees,
        "mode_frais_agence": mode_frais_agence,
        "taux_tva_frais_agence": vat_rate,
        "tva_frais_agence": vat,
        "total_ht": total_ht,
        "total_ttc": total,
    }


# Shared palette and wording from the supplied FPC Consulting workbook.
BUDGET_TITLE = "FFDE66"
BUDGET_SECTION = "FFC000"
BUDGET_HEADER = "E7E6E6"
BUDGET_SUBTOTAL = "AEAAAA"
BUDGET_IMPLEMENTATION = "C6E0B4"
BUDGET_TOTAL = "FFFF00"
BUDGET_NOTES = "ED7D31"
BUDGET_NOTICE = "A9D18E"
BUDGET_DISCLAIMER = (
    "Ce document est un estimatif budgétaire et non un devis. Il est par conséquent "
    "amené à être revu à la hausse ou à la baisse, en fonction des options et "
    "caractéristiques définitivement retenues."
)
AGENCY_DETAILS = (
    "Honoraires de conception, d'organisation, de mise en œuvre opérationnelle "
    "et d'animations d'activités\n"
    "- Matériels et port\n"
    "- Ressources humaines\n"
    "- Transport de l'équipe\n"
    "- Restauration de l'équipe\n"
    "- Hébergement de l'équipe"
)


def _document_totals(data: dict[str, Any]) -> dict[str, Any]:
    required = ("reference", "client", "nombre_personnes", "date_budget", "titre")
    missing = [field for field in required if not data.get(field)]
    if missing:
        raise ValueError(f"Champs obligatoires manquants: {', '.join(missing)}")
    if int(data["nombre_personnes"]) <= 0:
        raise ValueError("nombre_personnes doit être supérieur à zéro.")
    totals = calculate_budget_totals(
        data.get("sections") or [], data.get("frais_agence") or 0,
        data.get("taux_tva_frais_agence") or 0, data.get("mode_frais_agence") or "montant",
    )
    if not totals["sections"]:
        raise ValueError("Le budget doit contenir au moins une ligne.")
    return add_persisted_taxes(totals, data)


def _display_sections(totals: dict[str, Any]) -> list[dict[str, Any]]:
    """Show persisted taxes before fees without including them in the agency fee base."""
    sections = list(totals["sections"])
    if totals["taxes_selectionnees"]:
        lines = []
        for tax in totals["taxes_selectionnees"]:
            amount = tax["montant"]
            per_night = tax.get("tarif") is not None
            label = str(tax.get("nom") or tax.get("code") or "Taxe")
            if tax.get("taux") is not None:
                label += f" ({_format_quantity(tax['taux'])} %)"
            lines.append({
                "designation": label,
                "quantite": tax.get("nombre_chambres", 1) if per_night else 1,
                "prix_unitaire": tax["tarif"] if per_night else amount,
                "nombre_jours": tax.get("nombre_nuitees", 1) if per_night else 1,
                "montant_ht": amount,
                # Fiscal amounts are snapshots, not editable recalculations of current rates.
                "persisted_tax": True,
            })
        sections.append({
            "nom": "Taxes et contributions",
            "prestations": lines, "sous_total": totals["total_taxes"], "persisted_tax": True,
        })
    return sections


def _document_money(value: Any) -> str:
    amount = Decimal(str(value or 0))
    decimals = 0 if amount == amount.to_integral_value() else 2
    return f"{amount:,.{decimals}f}".replace(",", " ").replace(".", ",") + " CFA"


def is_budget_pdf_current(path: Path) -> bool:
    if not path.is_file():
        return False
    dependencies = (Path(__file__), ASSETS_DIR / "logo.png")
    return path.stat().st_mtime >= max(item.stat().st_mtime for item in dependencies if item.is_file())


def _draw_page(canvas: Any, document: Any) -> None:
    canvas.saveState()
    width, _ = A4
    canvas.setFont("Helvetica", 6)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(10 * mm, 7 * mm, "Ivoir Trips International — Budget estimatif")
    canvas.drawRightString(width - 10 * mm, 7 * mm, f"Page {document.page}")
    canvas.restoreState()


def generate_budget_pdf(data: dict[str, Any], output_dir: str | Path | None = None) -> str:
    totals = _document_totals(data)
    output_path = _safe_document_path(data["reference"], "pdf", output_dir)
    width = 190 * mm
    columns = [width * ratio for ratio in (.39, .12, .195, .12, .175)]
    styles = _styles()
    styles["section"].alignment = TA_CENTER
    for key in ("normal", "center", "right", "right_bold", "section"):
        styles[key].fontSize = 7
        styles[key].leading = 8.7
    styles["title"].fontSize = 10
    styles["title"].leading = 13
    notice_style = ParagraphStyle(
        "BudgetNotice", parent=styles["center"], fontName="Helvetica-Bold",
        fontSize=9, leading=12,
    )
    vat_style = ParagraphStyle(
        "BudgetVat", parent=styles["right_bold"], textColor=colors.red,
    )
    document = SimpleDocTemplate(
        str(output_path), pagesize=A4,
        leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=9 * mm, bottomMargin=13 * mm,
        title=f"Budget estimatif {data['reference']}", author="Ivoir Trips International",
    )
    story: list[Any] = []
    logo_path = _first_existing_asset("logo.png")
    if logo_path:
        logo_width, logo_height = _image_dimensions(logo_path, 76 * mm, 24 * mm)
        logo = Image(str(logo_path), width=logo_width, height=logo_height, hAlign="CENTER")
        story.extend([logo, Spacer(1, 2 * mm)])
    title = Table([
        [_paragraph("BUDGET ESTIMATIF", styles["title"])],
        [_paragraph(str(data["titre"]).upper(), styles["section"])],
    ], colWidths=[width])
    title.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + BUDGET_TITLE)),
        ("BOX", (0, 0), (-1, -1), .7, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(title)
    header = Table([
        ["DATE :", _display_date(data["date_budget"]), "", "LIEU :", _budget_location(data)],
        ["CLIENT :", str(data["client"]).upper(), "DURÉE", "JOUR(S) :", data.get("duree_jours") or 1],
        ["NOMBRE DE PARTICIPANTS :", data["nombre_personnes"], "", "NUIT(S) :", _budget_nights(data)],
    ], colWidths=columns)
    for row in header._cellvalues:
        for i, value in enumerate(row):
            row[i] = _paragraph(value, styles["normal"] if i in (0, 3) else styles["center"])
    header.setStyle(TableStyle([
        ("GRID", (0, 0), (1, -1), .6, colors.black),
        ("GRID", (3, 0), (-1, -1), .6, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.extend([header, Spacer(1, 6 * mm)])
    rows = [[_paragraph(label, styles["section"]) for label in (
        "Désignation", "Quantité", "Coût unitaire", "Nombre de jours", "Coût total",
    )]]
    commands = [
        ("GRID", (0, 0), (-1, -1), .6, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BUDGET_HEADER)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for section in _display_sections(totals):
        first = len(rows)
        rows.append([_paragraph(section["nom"].upper(), styles["section"]), "", "", "", ""])
        commands.extend([
            ("SPAN", (0, first), (-1, first)),
            ("BACKGROUND", (0, first), (-1, first), colors.HexColor("#" + BUDGET_SECTION)),
            ("NOSPLIT", (0, first), (-1, first + 1)),
        ])
        for item in section["prestations"]:
            rows.append([
                _paragraph(item["designation"], styles["normal"]),
                _paragraph(_format_quantity(item["quantite"]), styles["center"]),
                _paragraph(_document_money(item["prix_unitaire"]), styles["right"]),
                _paragraph(_format_quantity(item["nombre_jours"]), styles["center"]),
                _paragraph(_document_money(item["montant_ht"]), styles["right"]),
            ])
        subtotal = len(rows)
        rows.append([
            _paragraph(f"SOUS-TOTAL {section['nom'].upper()}", styles["section"]), "", "", "",
            _paragraph(_document_money(section["sous_total"]), styles["right_bold"]),
        ])
        commands.extend([
            ("SPAN", (0, subtotal), (3, subtotal)),
            ("BACKGROUND", (0, subtotal), (-1, subtotal), colors.HexColor("#" + BUDGET_SUBTOTAL)),
            ("NOSPLIT", (0, subtotal - 1), (-1, subtotal)),
        ])
        # Small rubrics should not be cut between pages (notably the two hotel taxes).
        if len(section["prestations"]) <= 3:
            commands.append(("NOSPLIT", (0, first), (-1, subtotal)))
        spacer = len(rows)
        rows.append([""] * 5)
        commands.extend([
            ("SPAN", (0, spacer), (-1, spacer)),
            ("TOPPADDING", (0, spacer), (-1, spacer), 1),
            ("BOTTOMPADDING", (0, spacer), (-1, spacer), 1),
        ])
    row = len(rows)
    rows.append([
        _paragraph("SOUS-TOTAL MISE EN ŒUVRE", styles["section"]), "", "", "",
        _paragraph(_document_money(totals["sous_total_ht"] + totals["total_taxes"]), styles["right_bold"]),
    ])
    commands.extend([
        ("SPAN", (0, row), (3, row)),
        ("BACKGROUND", (0, row), (-1, row), colors.HexColor("#" + BUDGET_IMPLEMENTATION)),
        ("NOSPLIT", (0, row - 2), (-1, row)),
    ])
    table = Table(rows, colWidths=columns, repeatRows=1, hAlign="CENTER")
    table.setStyle(TableStyle(commands))
    story.extend([table, Spacer(1, 4 * mm)])

    fee_title = "HONORAIRES D'AGENCE"
    if totals["mode_frais_agence"] == "pourcentage":
        fee_title += " — 17,5 % des prestations HT"
    vat_rate = _format_quantity(totals["taux_tva_frais_agence"])
    footer_rows = [
        [_paragraph(fee_title, styles["section"]), "", "", "", ""],
        [Paragraph(escape(AGENCY_DETAILS).replace("\n", "<br/>"), styles["normal"]), "", "", "",
         _paragraph(_document_money(totals["frais_agence"]), styles["right_bold"])],
        [_paragraph("TOTAL AVANT TVA SUR HONORAIRES", styles["section"]), "", "", "",
         _paragraph(_document_money(totals["total_ht"] + totals["total_taxes"]), styles["right_bold"])],
        ["", "", _paragraph(f"TVA {vat_rate} %", vat_style), "",
         _paragraph(_document_money(totals["tva_frais_agence"]), vat_style)],
        ["", "", _paragraph("TOTAL TTC", styles["section"]), "",
         _paragraph(_document_money(totals["total_ttc"]), styles["right_bold"])],
        ["", "", _paragraph("TOTAL TTC / PERSONNE", styles["section"]), "",
         _paragraph(_document_money((totals["total_ttc"] / Decimal(str(data["nombre_personnes"]))).quantize(Decimal("1"), rounding=ROUND_HALF_UP)), styles["right_bold"])],
    ]
    notes = _budget_notes(data)
    note_style = ParagraphStyle("BudgetNotes", parent=styles["normal"], fontName="Helvetica-Bold")
    for index, note in enumerate(notes):
        row = 4 + index
        while len(footer_rows) <= row:
            footer_rows.append([""] * 5)
        footer_rows[row][0] = _paragraph(note, note_style)
    footer = Table(footer_rows, colWidths=columns)
    footer_commands = [
        ("SPAN", (0, 0), (-1, 0)),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BUDGET_SECTION)),
        ("SPAN", (0, 1), (2, 1)), ("SPAN", (0, 2), (3, 2)),
        ("GRID", (0, 0), (-1, 2), .6, colors.black),
        ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#" + BUDGET_SUBTOTAL)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("VALIGN", (4, 1), (4, 1), "BOTTOM"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for row in (3, 4, 5):
        footer_commands.extend([
            ("SPAN", (2, row), (3, row)),
            ("GRID", (2, row), (-1, row), .6, colors.black),
            ("BACKGROUND", (2, row), (3 if row == 3 else 4, row),
             colors.HexColor("#" + (BUDGET_SUBTOTAL if row == 3 else BUDGET_TOTAL))),
        ])
    for row in range(4, 4 + len(notes)):
        footer_commands.extend([
            ("SPAN", (0, row), (1, row)),
            ("BACKGROUND", (0, row), (1, row), colors.HexColor("#" + BUDGET_NOTES)),
        ])
    footer.setStyle(TableStyle(footer_commands))
    notice = Table([[_paragraph(BUDGET_DISCLAIMER, notice_style)]], colWidths=[width - 8 * mm])
    notice.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + BUDGET_NOTICE)),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#4472C4")),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.extend([KeepTogether(footer), Spacer(1, 4 * mm), notice])
    document.build(story, onFirstPage=_draw_page, onLaterPages=_draw_page)
    return str(output_path)


def _append_budget_sheet(workbook: Workbook, data: dict[str, Any]) -> dict[str, Any]:
    totals = _document_totals(data)
    base_title = _safe_sheet_title(data.get("site_nom") or data.get("titre") or "Budget")
    title, suffix = base_title, 2
    while title.casefold() in {name.casefold() for name in workbook.sheetnames}:
        marker = f" ({suffix})"
        title = base_title[:31 - len(marker)] + marker
        suffix += 1
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 60
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    # Same column proportions, typeface and worksheet scale as the reference.
    for column, size in {"A": 1.285, "B": 118.71, "C": 36.57, "D": 59.285, "E": 36.285, "F": 53.71}.items():
        sheet.column_dimensions[column].width = size
    edge = Side(style="thin", color="000000")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    money_format = '#,##0 "CFA"'
    decimal_money_format = '#,##0.00 "CFA"'

    def style_row(row, first=2, last=6, fill=None, bold=False, height=35):
        sheet.row_dimensions[row].height = height
        for col in range(first, last + 1):
            cell = sheet.cell(row, col)
            cell.font = Font(name="Agency FB", size=18, bold=bold, color="000000")
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)

    def merge(row, first, last, value, *, align="center"):
        sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        cell = sheet.cell(row, first, value)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        return cell

    def money(row, value):
        cell = sheet.cell(row, 6, value)
        cell.number_format = money_format
        cell.alignment = Alignment(horizontal="right", vertical="center")
        if isinstance(value, (int, float, Decimal)) and Decimal(str(value)) % 1:
            cell.number_format = decimal_money_format
        return cell

    sheet.row_dimensions[1].height = 6.4
    for row in range(2, 6):
        sheet.row_dimensions[row].height = 14.65
    sheet.row_dimensions[6].height = 105
    # Center the image using its actual aspect ratio within the full B:F width.
    logo_path = _first_existing_asset("logo.png")
    if logo_path:
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        logo = ExcelImage(str(logo_path))
        logo.width, logo.height = _image_dimensions(logo_path, 610, 175)
        table_pixels = sum(int(sheet.column_dimensions[c].width * 7 + 5) for c in "BCDEF")
        logo.anchor = OneCellAnchor(
            _from=AnchorMarker(col=1, row=1, colOff=pixels_to_EMU((table_pixels - logo.width) / 2), rowOff=pixels_to_EMU(12)),
            ext=XDRPositiveSize2D(pixels_to_EMU(logo.width), pixels_to_EMU(logo.height)),
        )
        sheet.add_image(logo)

    for row, value, height in ((7, "BUDGET ESTIMATIF", 63), (8, str(data["titre"]).upper(), 42.6)):
        style_row(row, fill=BUDGET_TITLE, bold=True, height=height)
        merge(row, 2, 6, value)
    participants = int(data["nombre_personnes"])
    header = {
        "B9": "DATE :", "C9": _display_date(data["date_budget"]),
        "E9": "LIEU :", "F9": _budget_location(data),
        "B10": "CLIENT :", "C10": str(data["client"]).upper(),
        "D10": "DURÉE", "E10": "JOUR(S) :", "F10": int(data.get("duree_jours") or 1),
        "B11": "NOMBRE DE PARTICIPANTS :", "C11": participants,
        "E11": "NUIT(S) :", "F11": _budget_nights(data),
    }
    for row in (9, 10, 11):
        style_row(row, bold=True, height=33)
        sheet.cell(row, 4).border = Border()
    for coordinate, value in header.items():
        cell = sheet[coordinate]
        cell.value = value
        cell.alignment = Alignment(
            horizontal="left" if coordinate[0] in "BE" else "center",
            vertical="center", wrap_text=True,
        )
    sheet["D10"].alignment = Alignment(horizontal="right", vertical="center")
    for row in (12, 13, 14):
        sheet.row_dimensions[row].height = 22

    style_row(15, fill=BUDGET_HEADER, bold=True, height=27)
    for col, label in enumerate(("Désignation", "Quantité", "Coût unitaire", "Nombre de jours", "Coût total"), 2):
        sheet.cell(15, col, label)
    current_row = 16
    subtotal_rows, service_subtotal_rows = [], []
    for section in _display_sections(totals):
        name = str(section["nom"]).upper()
        style_row(current_row, fill=BUDGET_SECTION, bold=True)
        merge(current_row, 2, 6, name)
        current_row += 1
        first_line_row = current_row
        for item in section["prestations"]:
            # Wrap long descriptions without clipping custom prestations.
            text_lines = sum(max(1, (len(line) + 84) // 85) for line in item["designation"].splitlines())
            style_row(current_row, height=max(35, text_lines * 24 + 10))
            sheet.cell(current_row, 2, item["designation"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            quantity = _excel_number(item.get("quantite"), 1)
            price = _excel_number(item.get("prix_unitaire"))
            days = _excel_number(item.get("nombre_jours"), 1)
            amount = Decimal(str(item["montant_ht"]))
            sheet.cell(current_row, 3, quantity)
            sheet.cell(current_row, 4, price).number_format = decimal_money_format if Decimal(str(price)) % 1 else money_format
            sheet.cell(current_row, 4).alignment = Alignment(horizontal="right", vertical="center")
            sheet.cell(current_row, 5, days)
            calculated = (Decimal(str(quantity)) * Decimal(str(price)) * Decimal(str(days))).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
            value = (
                f"=ROUND(C{current_row}*D{current_row}*E{current_row},0)"
                if not item.get("persisted_tax") and calculated == amount else _excel_number(amount)
            )
            money(current_row, value)
            current_row += 1
        style_row(current_row, fill=BUDGET_SUBTOTAL, bold=True)
        merge(current_row, 2, 5, f"SOUS-TOTAL {name}")
        money(current_row, f"=SUM(F{first_line_row}:F{current_row - 1})")
        subtotal_rows.append(current_row)
        if not section.get("persisted_tax"):
            service_subtotal_rows.append(current_row)
        sheet.row_dimensions[current_row + 1].height = 25
        current_row += 2

    implementation_row = current_row
    style_row(current_row, fill=BUDGET_IMPLEMENTATION, bold=True)
    merge(current_row, 2, 5, "SOUS-TOTAL MISE EN ŒUVRE")
    money(current_row, "=" + "+".join(f"F{row}" for row in subtotal_rows))
    sheet.row_dimensions[current_row + 1].height = 25
    current_row += 2

    agency_title_row = current_row
    style_row(current_row, fill=BUDGET_SECTION, bold=True, height=45)
    title = "HONORAIRES D'AGENCE"
    if totals["mode_frais_agence"] == "pourcentage":
        title += " — 17,5 % DES PRESTATIONS HT"
    merge(current_row, 2, 6, title)
    current_row += 1
    agency_amount_row = current_row
    style_row(current_row, height=202)
    merge(current_row, 2, 4, AGENCY_DETAILS, align="left").alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    fee_base = "+".join(f"F{row}" for row in service_subtotal_rows)
    money(current_row, (
        f"=ROUND(({fee_base})*17.5%,0)"
        if totals["mode_frais_agence"] == "pourcentage" else int(totals["frais_agence"])
    )).alignment = Alignment(horizontal="right", vertical="bottom")
    current_row += 1

    total_ht_row = current_row
    style_row(current_row, fill=BUDGET_SUBTOTAL, bold=True, height=43)
    merge(current_row, 2, 5, "TOTAL AVANT TVA SUR HONORAIRES")
    money(current_row, f"=F{implementation_row}+F{agency_amount_row}")
    current_row += 1

    vat_row = current_row
    vat_rate = _format_quantity(totals["taux_tva_frais_agence"])
    style_row(current_row, first=4, bold=True, height=46)
    merge(current_row, 4, 5, f"TVA {vat_rate} %").fill = PatternFill("solid", fgColor=BUDGET_SUBTOTAL)
    money(current_row, f"=ROUND(F{agency_amount_row}*{vat_rate}/100,0)")
    for col in (4, 6):
        sheet.cell(current_row, col).font = Font(name="Agency FB", size=18, bold=True, color="FF0000")
    current_row += 1
    total_ttc_row = current_row
    style_row(current_row, first=4, fill=BUDGET_TOTAL, bold=True, height=47)
    merge(current_row, 4, 5, "TOTAL TTC")
    money(current_row, f"=F{total_ht_row}+F{vat_row}")
    current_row += 1
    total_person_row = current_row
    style_row(current_row, first=4, fill=BUDGET_TOTAL, bold=True, height=35)
    merge(current_row, 4, 5, "TOTAL TTC / PERSONNE")
    money(current_row, f"=F{total_ttc_row}/C11")

    notes = _budget_notes(data)
    for index, note in enumerate(notes):
        row = total_ttc_row + index
        # Notes occupy B only, as in the reference, beside the summary in D:F.
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, 32 * max(1, (len(note) + 79) // 80))
        cell = sheet.cell(row, 2, note)
        cell.font = Font(name="Agency FB", size=18, bold=True)
        cell.fill = PatternFill("solid", fgColor=BUDGET_NOTES)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    disclaimer_row = max(total_person_row, total_ttc_row + len(notes) - 1) + 2
    style_row(disclaimer_row, fill=BUDGET_NOTICE, bold=True, height=180)
    cell = merge(disclaimer_row, 2, 6, BUDGET_DISCLAIMER)
    cell.font = Font(name="Calibri", size=24, bold=True, color="000000")
    notice_edge = Side(style="thin", color="4472C4")
    cell.border = Border(left=notice_edge, right=notice_edge, top=notice_edge, bottom=notice_edge)
    sheet.freeze_panes = None
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = "15:15"
    sheet.page_margins.left = sheet.page_margins.right = .3
    sheet.page_margins.top = sheet.page_margins.bottom = .4
    sheet.print_area = f"B1:F{disclaimer_row}"
    return {
        "sheet": sheet.title, "subtotal": implementation_row, "fees": agency_amount_row,
        "vat": vat_row, "total": total_ttc_row, "per_person": total_person_row,
    }


def generate_budget_excel(data: dict[str, Any], output_dir: str | Path | None = None) -> str:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_budget_sheet(workbook, data)
    output_path = _safe_document_path(data["reference"], "xlsx", output_dir)
    workbook.save(output_path)
    return str(output_path)


def generate_multi_site_budget_excel(
    budgets: list[dict[str, Any]], reference: str, output_dir: str | Path | None = None,
) -> str:
    if not budgets or len({data["offre_id"] for data in budgets}) != 1:
        raise ValueError("Le comparatif doit contenir les budgets d'une seule offre.")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Comparatif"
    summary.sheet_view.showGridLines = False
    summary.append(["Site", "Référence", "Statut", "Mise en œuvre (taxes incluses)", "Frais d'agence", "TVA sur honoraires", "Total TTC", "TTC / personne"])
    labels = {"valide": "Validé", "genere": "Généré", "brouillon": "Brouillon", "annule": "Annulé"}
    for data in budgets:
        cells = _append_budget_sheet(workbook, data)
        sheet_name = quote_sheetname(cells["sheet"])
        summary.append([
            data.get("site_nom") or data["titre"], data["reference"], labels.get(data.get("statut"), "Brouillon"),
            *[f"={sheet_name}!F{cells[key]}" for key in ("subtotal", "fees", "vat", "total", "per_person")],
        ])
        row = summary.max_row
        summary.cell(row, 1).hyperlink = f"#{sheet_name}!B9"
        for cell in summary[row]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.column >= 4:
                cell.number_format = '#,##0 "FCFA"'
            if data.get("statut") == "valide":
                cell.fill = PatternFill("solid", fgColor="EAF7F0")
        summary.row_dimensions[row].height = 36
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="008A50")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    summary.row_dimensions[1].height = 32
    for column in "ABCDEFGH":
        summary.column_dimensions[column].width = 32 if column == "A" else 22
    summary.freeze_panes = "D2"
    summary.auto_filter.ref = summary.dimensions
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    output_path = _safe_document_path(reference, "xlsx", output_dir)
    workbook.save(output_path)
    return str(output_path)
